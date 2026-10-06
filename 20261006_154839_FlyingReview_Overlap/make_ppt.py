from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pathlib import Path
import json

ev=json.loads(Path('out/Calculation_evidence.json').read_text());c=ev['fixed'];r=ev['defaultResult']
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
prs.core_properties.title='0선방어 Flying Review 검토자료';prs.core_properties.subject='Basler a2a2600-20gcBAS / Mitutoyo 378-810-3 · 서로 다른 Cell의 2/4 Head Review';prs.core_properties.author='A3 Laser Driller';prs.core_properties.keywords='Flying Vision, Zero Defence, Basler, Motion Blur'
FONT='Noto Sans CJK KR';NAVY='10243A';BLUE='1269D3';CYAN='009FA9';GRAY='5D6E80';LIGHT='F3F6FA';GREEN='0F805B';RED='BC3948';ORANGE='EFB757'
def shape(s,x,y,w,h,fill,line=None,kind=MSO_SHAPE.ROUNDED_RECTANGLE):
 q=s.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h));q.fill.solid();q.fill.fore_color.rgb=RGBColor.from_string(fill)
 if line:q.line.color.rgb=RGBColor.from_string(line);q.line.width=Pt(1)
 else:q.line.fill.background()
 if kind==MSO_SHAPE.ROUNDED_RECTANGLE:
  try:q.adjustments[0]=.10
  except:pass
 return q
def txt(s,x,y,w,h,text,size=20,color=NAVY,bold=False,align=PP_ALIGN.LEFT):
 q=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=q.text_frame;tf.clear();tf.word_wrap=True;tf.margin_left=Inches(.02);tf.margin_right=Inches(.02);tf.margin_top=Inches(.01);tf.margin_bottom=0
 for i,line in enumerate(text.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.alignment=align;p.space_before=Pt(0);p.space_after=Pt(4)
  for run in p.runs:run.font.name=FONT;run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor.from_string(color)
 return q
def ln(s,x1,y1,x2,y2,color=GRAY,width=2):
 q=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));q.line.color.rgb=RGBColor.from_string(color);q.line.width=Pt(width);return q
def slide(n,title,sub):
 s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(LIGHT)
 shape(s,0,0,13.3333,1.10,NAVY,kind=MSO_SHAPE.RECTANGLE)
 txt(s,.42,.15,12.45,.51,title,27,'FFFFFF',True);txt(s,.44,.72,12,.23,sub,11,'CFE2F4')
 txt(s,.44,7.08,11.8,.25,'A3 Laser Driller · 2026-10-06 · 명목 설계 계산 / 실제 장비·영상 검증 필요',10,GRAY)
 txt(s,12.32,7.02,.52,.30,f'{n}/3',12,GRAY,align=PP_ALIGN.RIGHT)
 return s

# Slide 1: common reverse motion, physical station overlap.
s=slide(1,'정방향 투입 후, 역방향 MOF에서 가공과 리뷰를 함께 진행','Scanner를 먼저 만나고 1467 mm 뒤 Review · 이미 가공된 점부터 측정')
for i,(title,body) in enumerate([('① 정방향 투입','가공대기 위치까지 이동'),('② 역방향 가공','Y 등속 · Scanner 먼저'),('③ 리뷰 병행','앞부분 리뷰 / 뒷부분 가공')]):
 x=.45+i*4.2;shape(s,x,1.34,4.0,1.04,'FFFFFF');txt(s,x+.16,1.46,3.7,.32,title,21,BLUE,True);txt(s,x+.16,1.95,3.7,.27,body,15)
shape(s,.45,2.63,12.4,2.24,'FFFFFF');txt(s,.67,2.79,11.9,.33,'기판이 두 위치를 동시에 걸치는 구간이 있습니다',21,NAVY,True)
# Geometry drawn to a common scale; 1500 mm board straddles stations 1467 mm apart.
x0=2.60;scale=6.55/1467;x1=x0+1467*scale
shape(s,x0-.57,3.36,1.14,.48,BLUE);txt(s,x0-.55,3.45,1.10,.26,'Scanner',16,'FFFFFF',True,PP_ALIGN.CENTER)
shape(s,x1-.57,3.36,1.14,.48,CYAN);txt(s,x1-.55,3.45,1.10,.26,'Review',16,'FFFFFF',True,PP_ALIGN.CENTER)
ln(s,x0,3.97,x1,3.97,ORANGE,2);txt(s,5.09,3.79,1.50,.27,'1467 mm',15,NAVY,True,PP_ALIGN.CENTER)
shape(s,x0-16.5*scale,4.21,1500*scale,.35,'D7EAFF',kind=MSO_SHAPE.RECTANGLE);txt(s,3.64,4.22,4.44,.25,'기판 1500 mm · 역방향 이동 →',16,BLUE,True,PP_ALIGN.CENTER)
ln(s,x0,3.84,x0,4.70,BLUE,1);ln(s,x1,3.84,x1,4.70,CYAN,1)
shape(s,.45,5.13,5.98,1.44,'EAF3FF');txt(s,.66,5.28,5.58,.35,'같은 점의 도착 지연',20,BLUE,True);txt(s,.66,5.80,5.58,.42,'1467 ÷ Y100 = 14.67초',25,NAVY,True)
shape(s,6.65,5.13,6.20,1.44,'DEF4F1');txt(s,6.86,5.28,5.78,.35,'기판 전체의 통과창 중첩',20,CYAN,True);txt(s,6.86,5.80,5.78,.42,'(1500 − 1467) ÷ 100 = 0.33초',24,NAVY,True)
txt(s,.56,6.77,12.05,.24,'통과창 중첩과 실제 Laser ON 중첩은 구분합니다. 실제 출사는 레시피 유효 가공 범위로 확인합니다.',12,GRAY)
s.notes_slide.notes_text_frame.text='사용자 동작: 투입후 +StageY 정방향으로 가공대기, -StageY 역방향 MOF에서 Scanner 먼저, 공통간격1467mm 뒤 Review. 전체가공완료 대기없이 해당측정점의가공후 촬영. 기판앞부분리뷰중뒷부분가공이가능한 컨셉이며 실제가공범위에따라 Shot동시성달라짐. 그림의 기판1500mm와 스테이션간격1467mm는 동일축척, 33mm물리통과창중첩. 현재 test0929는 유효가공Y16.5125~859.7125(843.2mm)로 실제Shot창이일찍끝남. 바닥이동의MOF통과창과실제LaserON을동일시하지않음. 정방향기판좌표는설계가정. 출처/상세수식은Formula Guide에포함.'

# Slide 2: exact reverse timeline and board quotas.
s=slide(2,'가공 종료 뒤에도 리뷰가 남는 시간을 같은 시간축에서 계산','기본값: Y100 / X200 mm/s · 노출2 µs · 허용blur2 px · 시간0 = 역방향 시작')
shape(s,.45,1.30,12.40,3.05,'FFFFFF');x0=2.83;sc=9.45/r['profile']['total'];xx=lambda t:x0+t*sc
bars=[('MOF 가공주행',r['processStart'],r['processEnd'],1.91,BLUE),('Review 기판 통과',r['reviewEnter'],r['reviewExit'],2.45,CYAN),('레시피 유효 Shot',r['firstTime'],r['lastTime'],2.99,'9CAFC4'),('4점 리뷰 작업',r['modes']['4']['groups'][0]['reviewStart'],r['modes']['4']['groups'][0]['reviewEnd'],3.53,CYAN)]
for t in range(0,36,5):
 ln(s,xx(t),1.78,xx(t),3.94,'DFE8EF',1);txt(s,xx(t)-.2,1.46,.43,.20,f'{t}s',11,GRAY,align=PP_ALIGN.CENTER)
for label,a,b,y,col in bars:
 txt(s,.66,y+.03,2.10,.27,label,15,NAVY,True);shape(s,xx(a),y,(b-a)*sc,.32,col,kind=MSO_SHAPE.RECTANGLE)
g=r['modes']['4']['groups'][0];ln(s,xx(r['processEnd']),1.82,xx(r['processEnd']),4.05,ORANGE,2)
for n in g['nodes']:
 ln(s,xx(n['cross']),3.47,xx(n['cross']),3.92,NAVY,1.5);txt(s,xx(n['cross'])-.2,3.98,.44,.20,f"H{n['head']:02}",10,NAVY,True,PP_ALIGN.CENTER)
txt(s,7.75,1.82,4.45,.25,'主행창 종료 20.100s'.replace('主','주'),12,NAVY,True,PP_ALIGN.RIGHT)
for x,w,title,value,detail in [(.45,5.98,'2개 스캐너 / 기판','4기판 · 10분','H01→02 / H03→04 / H05→06 / H07→08\n주행창 이후 잔여: 1.832 / 5.832 / 1.832 / 5.832초'),(6.65,6.20,'4개 스캐너 / 기판','2기판 · 5분','H01→02→03→04 / H05→06→07→08\n주행창 이후 잔여: 각5.832초')]:
 shape(s,x,4.60,w,1.56,'EAF3FF' if x<1 else 'DEF4F1');txt(s,x+.18,4.73,w-.36,.27,title,17,BLUE,True);txt(s,x+.18,5.08,w-.36,.37,value,23,NAVY,True);txt(s,x+.18,5.54,w-.36,.53,detail,12,NAVY)
txt(s,.56,6.37,12.05,.49,'현재 test0929: 실제 Shot창 5.265~13.697s / 첫 리뷰 19.932s → 실제 Shot–리뷰 중첩 0s\n유효 가공Y범위843.2mm 기준. 기판 통과창0.33s 중첩과 별도이며, 실제 연속 가공 범위가 바뀌면 재산정합니다.',12,GRAY)
s.notes_slide.notes_text_frame.text=f'시간0=역방향시작. MOF기판스캐너통과 {r["processStart"]:.6f}~{r["processEnd"]:.6f}, Review기판통과 {r["reviewEnter"]:.6f}~{r["reviewExit"]:.6f}, 명목Shot {r["firstTime"]:.6f}~{r["lastTime"]:.6f}. 4점리뷰작업 {g["reviewStart"]:.6f}~{g["reviewEnd"]:.6f}, MOF주행창중첩 {g["overlap"]:.6f}, 주행창종료후촬영잔여 {g["reviewTail"]:.6f}. 마지막명목Shot종료기준잔여 {g["shotTail"]:.6f}초. 실제ProcessDone은Scanner/Laser완료ACK이며명목마지막좌표와다를수있음. 작업창에는점사이X이동/정착/대기포함, 연속노출아님. 150초/기판계획Tact기준4기판600초,2기판300초. 각Head1점 Cell1,11,21,32,5,15,25,36 모두다름. 고정점을속도만바꿔종료순서가바뀌는것아님. 소유Head/Mask유효Shot검증포인트 유지.'

# Slide 3: independent completion, remaining exposure and quality gates.
s=slide(3,'가공완료와 리뷰완료는 독립적입니다','리뷰 시작은 해당 점의 가공 완료 후 · 전체 ProcessDone을 기다리지 않음')
shape(s,.45,1.34,12.40,.74,'FFFFFF');txt(s,.65,1.53,12.0,.34,'역방향 Y 등속 유지  →  해당 점 도착 시 촬영  →  실제 Frame 수신·Vision 판정',21,BLUE,True,PP_ALIGN.CENTER)
for x,w,title,body,col in [(.45,5.98,'A  가공 완료 시 리뷰도 완료','마지막 노출이 이미 끝남\nY 감속 가능 → 양쪽 완료신호 확인',GREEN),(6.65,6.20,'B  가공 완료 시 리뷰가 남음','남은 마지막 촬영까지 Y 등속 유지\n노출 종료 후 감속 → 양쪽 완료신호 확인',BLUE)]:
 shape(s,x,2.34,w,1.78,'DDF3EA' if x<1 else 'EAF3FF');txt(s,x+.18,2.52,w-.36,.37,title,20,col,True);txt(s,x+.18,3.10,w-.36,.75,body,18,NAVY)
shape(s,.45,4.38,12.40,.63,NAVY);txt(s,.67,4.54,11.97,.30,'ProcessDone  AND  ReviewDone(요구수량의 Frame·Vision 결과)  →  검사 완료',19,'FFFFFF',True,PP_ALIGN.CENTER)
for x,w,title,val,small in [(.45,3.97,'이동 blur','1.6px ≤2px','Y100 × 노출2µs ÷125'),(4.65,3.98,'X 선도착 최소 여유','0.730초','다음 점2초 ≥ 준비1.270초'),(8.86,3.99,'왕복 포함 계산 Tact','132.740초 / 150초','정방향35.870 + 역방향35.870 +61')]:
 shape(s,x,5.29,w,1.20,'FFFFFF');txt(s,x+.15,5.42,w-.30,.25,title,15,BLUE,True);txt(s,x+.15,5.81,w-.30,.35,val,21,NAVY,True);txt(s,x+.15,6.24,w-.30,.20,small,10,GRAY)
txt(s,.56,6.69,12.04,.25,'Basler a2a2600-20gcBAS + Mitutoyo378-810-3 SL20x · 전체8점은 다른 Cell · 입력8개 + 2/4개 선택',12,GRAY)
s.notes_slide.notes_text_frame.text='두종료경우는개념관계예시로현재고정레시피는B. 실제ProcessDone은Scanner/Laser작업완료ACK, ReviewDone은실제요구수량2/4 Frame수신및Vision결과확인. 노출끝은Frame/Vision완료아님. 부적합/미수신이면검사완료로처리하지않음. 마지막노출±5µs까지Y등속유지후감속가능. 촬영/판정중첩을자동Offset적용으로확대하지않음. 총배율20x=200mm튜브렌즈+1xRelay, 물체pixel2.5/20=.125µm. LensNA.28/WD30.5/R약1µm@550nm. BaslerGlobalShutter최소2µs18.5fps, Color. 허용1px이면Y≤62.5blur조건충족하나왕복Tact초과. X정착50ms Guard20ms 지연50µs Jitter±5µs 고정설계값. 기본최악H03→H04 ΔX200 ΔY200 X200 a=d1000, X이동1.2s 필요1.270062s 여유.729938s. Stage기판시작끝−1584.5mm, 가공대기1982.5mm: 각거리3567mm/Stroke3960, 정방향에도동일Y입력가정. 기타60초는계산한왕복이동제외, 영상여유1초는실제ReviewDone대체안됨. 공식출처 https://docs.baslerweb.com/a2a2600-20gcbas ; https://docs.baslerweb.com/exposure-time ; https://www.mitutoyo.com/webfoo/wp-content/uploads/E4191-378_010611.pdf p20/29. 확인2026-10-06. 가공15초/리뷰15초를별도직렬로더하지않는다. FollowError/조명/SNR/Scanner서비스시간미계산.'

out=Path('out/0선방어_검토자료.pptx');prs.save(out)
assert len(prs.slides)==3
for s in prs.slides:
 for q in s.shapes:
  assert q.left>=0 and q.top>=0 and q.left+q.width<=prs.slide_width+Inches(.02) and q.top+q.height<=prs.slide_height+Inches(.02)
print(f'Created {out}: 3 slides, editable text/shapes, all objects in slide bounds')
