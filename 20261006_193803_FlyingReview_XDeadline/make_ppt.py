from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pathlib import Path
import json

ev=json.loads(Path('out/Calculation_evidence.json').read_text());c=ev['fixed'];r=ev['inquiryResult']
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
prs.core_properties.title='0선방어 Flying Review 검토자료';prs.core_properties.subject='Basler a2a2600-20gcBAS / Mitutoyo 378-810-3 · 서로 다른 Cell의 2/4 Head Review';prs.core_properties.author='A3 Laser Driller';prs.core_properties.keywords='Flying Vision, Zero Defence, Basler, Motion Blur'
FONT='Noto Sans CJK KR';NAVY='10243A';BLUE='1269D3';CYAN='009FA9';GRAY='5D6E80';LIGHT='F3F6FA';GREEN='0F805B';RED='BC3948';ORANGE='EFB757'
def shape(s,x,y,w,h,fill,line=None,kind=MSO_SHAPE.ROUNDED_RECTANGLE):
 q=s.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h));q.fill.solid();q.fill.fore_color.rgb=RGBColor.from_string(fill)
 if line:q.line.color.rgb=RGBColor.from_string(line);q.line.width=Pt(1)
 else:q.line.fill.background()
 if kind==MSO_SHAPE.ROUNDED_RECTANGLE:
  try:q.adjustments[0]=.10
  except:pass
 return q
def txt(s,x,y,w,h,text,size=20,color=NAVY,bold=False,align=PP_ALIGN.LEFT):
 q=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=q.text_frame;tf.clear();tf.word_wrap=True;tf.margin_left=Inches(.02);tf.margin_right=Inches(.02);tf.margin_top=Inches(.01);tf.margin_bottom=0
 for i,line in enumerate(text.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.alignment=align;p.space_before=Pt(0);p.space_after=Pt(4)
  for run in p.runs:run.font.name=FONT;run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor.from_string(color)
 return q
def ln(s,x1,y1,x2,y2,color=GRAY,width=2):
 q=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));q.line.color.rgb=RGBColor.from_string(color);q.line.width=Pt(width);return q
def slide(n,title,sub):
 s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(LIGHT)
 shape(s,0,0,13.3333,1.10,NAVY,kind=MSO_SHAPE.RECTANGLE)
 txt(s,.42,.15,12.45,.51,title,27,'FFFFFF',True);txt(s,.44,.72,12,.23,sub,11,'CFE2F4')
 txt(s,.44,7.08,11.8,.25,'A3 Laser Driller · 2026-10-06 · 명목 설계 계산 / 실제 장비·영상 검증 필요',10,GRAY)
 txt(s,12.32,7.02,.52,.30,f'{n}/3',12,GRAY,align=PP_ALIGN.RIGHT)
 return s

# Slide 1: independent conditions.
s=slide(1,'27.667초는 통과합니다. 막히는 것은 촬영점 사이의2초입니다','문의 조건: Stroke2700mm · 노출10µs · 허용blur8px · X200/A=D150 · Y100/A=D150')
items=[('역방향30초','27.667초 ≤30초','가속0.667 + 등속26.333 + 감속0.667',GREEN),('Y 등속 후보점','4/4점이 등속 구간 안','등속거리2633.3mm · 기판1500mm 충분',GREEN),('이동 blur','8.00px = 허용8.00px','Y100 × 노출10µs →1.000µm 이동',GREEN),('X 준비·취득 간격','H03→H04에서0.379초 부족','다음 점2.000초 / 준비2.379초',RED)]
for i,(title,value,note,col) in enumerate(items):
 x=.45+(i%2)*6.2;y=1.39+(i//2)*1.77
 shape(s,x,y,6.0,1.48,'FFFFFF');shape(s,x,y,.08,1.48,col,kind=MSO_SHAPE.RECTANGLE)
 txt(s,x+.19,y+.15,5.6,.28,title,17,col,True);txt(s,x+.19,y+.54,5.6,.40,value,24,NAVY,True);txt(s,x+.19,y+1.08,5.6,.28,note,13,GRAY)
shape(s,.45,5.21,12.4,1.44,NAVY);txt(s,.68,5.38,11.96,.42,'전체 Stroke 종료보다, 각 촬영점의 도착시간이 더 빠른 마감입니다',24,'FFFFFF',True,PP_ALIGN.CENTER)
txt(s,.68,6.00,11.96,.35,'기존 “Stroke 내 촬영4/4” → “Y 등속 후보4/4”로 수정 · 촬영 완료를 보장하지 않습니다',17,'CFE2F4',align=PP_ALIGN.CENTER)
s.notes_slide.notes_text_frame.text='문의 조건 고정 선정점 기준. 실제 기판1500mm를 등속 가공할 거리는 충분하다. Y 가속거리33.3333, 감속거리33.3333, 등속거리2633.3333mm; cruise0.666667~27s. 모든 대표점 Y후보 범위 통과. 현재 조건 불가는 H3->4/H7->8 사이 X준비 마감 미달. 다른 유효점 선정 시 명목상 가능하며 장비 자체의 불가능 판정이 아니다. Frame/Vision 확인은 별도.'

# Slide 2: precise between-shot deadline.
s=slide(2,'H03 촬영 후2초, H04가 먼저 지나가는데 X는 아직 준비 중입니다','현재 고정점: ΔX200mm / ΔY200mm · Y100mm/s → 다음 점까지2.000초')
shape(s,.45,1.36,12.4,3.61,'FFFFFF');xx=lambda t:2.61+t*3.53
for t in [0,1,2]:
 ln(s,xx(t),1.89,xx(t),4.27,'E0E9F1',1);txt(s,xx(t)-.23,1.57,.56,.26,f'{t:.0f}s',13,GRAY,align=PP_ALIGN.CENTER)
for label,y in [('H04 도착',2.03),('X 준비',2.94),('카메라 주기',3.98)]:txt(s,.66,y,1.76,.34,label,19,NAVY,True)
shape(s,xx(0),2.03,2*3.53,.38,BLUE,kind=MSO_SHAPE.RECTANGLE);txt(s,xx(2)+.10,2.08,2.06,.32,'2.000초 마감',18,BLUE,True)
shape(s,xx(0),2.94,2.309401*3.53,.42,CYAN,kind=MSO_SHAPE.RECTANGLE);shape(s,xx(2.309401),2.94,.07007*3.53,.42,ORANGE,kind=MSO_SHAPE.RECTANGLE)
txt(s,4.02,3.00,5.10,.31,'가속·감속 이동2.309초',18,'FFFFFF',True,PP_ALIGN.CENTER);txt(s,9.58,3.47,2.91,.28,'정착·여유 포함2.379초',15,GRAY,align=PP_ALIGN.RIGHT)
shape(s,xx(0),3.98,.054064*3.53,.28,'9CAFC4',kind=MSO_SHAPE.RECTANGLE);txt(s,xx(.054064)+.14,3.98,7.08,.34,'0.054초 → 충분함 · 현재 실패 원인은 X 준비',17,GRAY)
ln(s,xx(2),1.89,xx(2),4.54,RED,2.5);ln(s,xx(2),4.54,xx(2.379471),4.54,RED,5);txt(s,8.80,4.61,3.58,.28,'0.379초 부족',18,RED,True,PP_ALIGN.RIGHT)
shape(s,.45,5.20,12.4,1.46,'FFE9EC');txt(s,.67,5.34,11.96,.38,'H04 마감26.165초 < 준비 예산 종료26.545초 < Stroke 종료27.667초',22,RED,True,PP_ALIGN.CENTER)
txt(s,.67,5.91,11.96,.51,'전체30초 안이어도 늦습니다. 목표점은 그 사이 Y37.947mm 더 이동합니다.\nH07→H08도2.000초 대비2.184초가 필요해0.184초 부족합니다.',16,NAVY,align=PP_ALIGN.CENTER)
s.notes_slide.notes_text_frame.text='X200mm, Vmax200, A=D150: 삼각형, Vpeak173.205mm/s, Tmove2sqrt(200/150)=2.309401s. 필요=max(Tmove+노출10us+정착50ms+Guard20ms+Trigger delay50us+2jitter5us, 1/18.5+2jitter)=2.379471s. H3노출시작24.165078, H4노출시작26.165078; 준비 예산26.544549. 0.379471s*100mm/s=37.9471mm. 렌즈시야Y0.266mm이므로 늦게 준비하고 동일점 촬영할 수 없다. 고정값 정착과Guard는 설계 가정. X 실제 S-Curve/Follow Error 확인 필요.'

# Slide 3: remedies are explicit comparisons.
s=slide(3,'현재 선정점이 문제입니다. 다른 유효점 또는 X 가감속으로 비교합니다','입력조건은 그대로 유지 가능 · 8개 Head의 Cell 모두 다름 · 비교안을 자동 적용하지 않음')
shape(s,.45,1.37,7.04,3.88,'FFFFFF');txt(s,.67,1.56,6.56,.35,'A. 측정점을 더 늦게 도착하는 유효점으로',21,BLUE,True)
txt(s,.67,2.20,6.56,.48,'H04: 같은Cell32 안에서 Y40.5mm 늦게',19,NAVY,True);txt(s,.67,2.75,6.56,.43,'도착2.405초 − 준비2.379초 = 여유25.5ms',18,GREEN,True)
txt(s,.67,3.39,6.56,.48,'H08: 같은Cell36 안에서 Y21.6mm 늦게',19,NAVY,True);txt(s,.67,3.94,6.56,.43,'도착2.216초 − 준비2.184초 = 여유31.9ms',18,GREEN,True)
txt(s,.67,4.59,6.56,.30,'X 이동거리 동일 · Stroke2700mm 내 Y 후보 유지',14,GRAY)
shape(s,7.70,1.37,5.15,3.88,'DEF4F1');txt(s,7.91,1.56,4.69,.35,'B. X 가감속250 비교',21,CYAN,True)
txt(s,7.91,2.32,4.69,.40,'X 속도200mm/s는 유지',18,NAVY,True);txt(s,7.91,2.95,4.69,.70,'H03→H04 준비1.870초\n2초 대비 여유0.130초',22,GREEN,True)
txt(s,7.91,4.06,4.69,.65,'현재 선정점 그대로 계산상 가능\n장비가 이 가감속을 허용하는지 확인',15,GRAY)
shape(s,.45,5.48,12.4,1.18,NAVY);txt(s,.68,5.65,11.96,.33,'시뮬레이터에서 입력을 바꾸면, 가이드의 수식·숫자·도식도 함께 바뀝니다',20,'FFFFFF',True,PP_ALIGN.CENTER)
txt(s,.68,6.13,11.96,.30,'A안: 고정 정착50ms·Guard20ms를 포함한 명목 여유 · 실제 위치·영상·정착 검증 필요',14,'CFE2F4',align=PP_ALIGN.CENTER)
s.notes_slide.notes_text_frame.text='A안: H4 Cell32 col2 row16 B7 targetQ-641.175 reviewX404.05; H8 Cell36 col2 row9 B7 targetQ-622.275 reviewX804.05. 고정 다른8Cell 유지, Shot/mask/head 소유권 geometryValid 확인. H4 gap2.405 need2.379471 slack25.5289ms, H8 gap2.216 need2.184150 slack31.8502ms. 최소Stroke2657.009333mm <2700. 계산상 두기판모두가능. 선택점예시만 표시하며 자동대체하지 않음. B안 X A=D250에서는 D200 v200 trapezoid Tmove1.8s+70.07ms=1.87007s; slack129.93ms. 동일X속도200에서 최소대칭가감속215.06995mm/s²이며 실제설비가능값 별도. Y84는 X간격 개선하지만 fullStroke32.702857s로30초실패하므로 단독해결아님. Basler2.5um/총20x=.125um/px, minExposure2us, 기본18.5fps; actualframe/vision processing 미계산.'

out=Path('out/0선방어_검토자료.pptx');prs.save(out);assert len(prs.slides)==3
for s in prs.slides:
 for q in s.shapes:
  assert q.left>=0 and q.top>=0 and q.left+q.width<=prs.slide_width+Inches(.02) and q.top+q.height<=prs.slide_height+Inches(.02)
print(f'Created {out}: 3 editable slides, all shapes within slide bounds')
