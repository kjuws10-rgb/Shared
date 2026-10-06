from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pathlib import Path
import json

ev=json.loads(Path('out/Calculation_evidence.json').read_text());c=ev['fixed'];r=ev['defaultResult']
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
prs.core_properties.title='0선방어 Flying Review 검토자료';prs.core_properties.subject='Basler a2a2600-20gcBAS / Mitutoyo 378-810-3 · 서로 다른 Cell의 2/4 Head Review';prs.core_properties.author='A3 Laser Driller';prs.core_properties.keywords='Flying Vision, Zero Defence, Basler, Motion Blur'
FONT='Noto Sans CJK KR';NAVY='10243A';BLUE='1269D3';CYAN='009FA9';GRAY='5D6E80';LIGHT='F3F6FA';GREEN='0F805B';RED='BC3948';ORANGE='EFB757'
def shape(s,x,y,w,h,fill,line=None,kind=MSO_SHAPE.ROUNDED_RECTANGLE):
 q=s.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h));q.fill.solid();q.fill.fore_color.rgb=RGBColor.from_string(fill)
 if line:q.line.color.rgb=RGBColor.from_string(line);q.line.width=Pt(1)
 else:q.line.fill.background()
 if kind==MSO_SHAPE.ROUNDED_RECTANGLE:
  try:q.adjustments[0]=.10
  except:pass
 return q
def txt(s,x,y,w,h,text,size=20,color=NAVY,bold=False,align=PP_ALIGN.LEFT):
 q=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=q.text_frame;tf.clear();tf.word_wrap=True;tf.margin_left=Inches(.02);tf.margin_right=Inches(.02);tf.margin_top=Inches(.01);tf.margin_bottom=0
 for i,line in enumerate(text.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.alignment=align;p.space_before=Pt(0);p.space_after=Pt(4)
  for run in p.runs:run.font.name=FONT;run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor.from_string(color)
 return q
def ln(s,x1,y1,x2,y2,color=GRAY,width=2):
 q=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));q.line.color.rgb=RGBColor.from_string(color);q.line.width=Pt(width);return q
def slide(n,title,sub):
 s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(LIGHT)
 shape(s,0,0,13.3333,1.10,NAVY,kind=MSO_SHAPE.RECTANGLE)
 txt(s,.42,.15,12.45,.51,title,27,'FFFFFF',True);txt(s,.44,.72,12,.23,sub,11,'CFE2F4')
 txt(s,.44,7.08,11.8,.25,'A3 Laser Driller · 2026-10-06 · 명목 설계 계산 / 실제 장비·영상 검증 필요',10,GRAY)
 txt(s,12.32,7.02,.52,.30,f'{n}/3',12,GRAY,align=PP_ALIGN.RIGHT)
 return s

# Slide 1: exposure and sampling.
s=slide(1,'움직이면서 촬영할 때, 노출이 blur를 결정합니다','Basler a2a2600-20gcBAS + Mitutoyo 378-810-3 M Plan Apo SL20x')
for x,w,title,body in [( .45,4.02,'카메라','2.5 µm pixel · Global Shutter\n최소 노출 2 µs · 기본 18.5 fps'),(4.65,4.02,'렌즈','총배율 20x 가정 · NA 0.28\nWD 30.5 mm · 200 mm 튜브렌즈'),(8.85,4.02,'물체면 환산','1 pixel = 0.125 µm\nFOV 0.325 × 0.266 mm')]:
 shape(s,x,1.34,w,1.22,'FFFFFF');txt(s,x+.16,1.43,w-.3,.3,title,15,BLUE,True);txt(s,x+.16,1.84,w-.3,.58,body,16)
shape(s,.45,2.80,6.11,3.48,'FFFFFF');txt(s,.67,2.96,5.70,.42,'Y100 mm/s · 노출2 µs',23,NAVY,True)
txt(s,.70,3.55,2.2,.3,'정지한 점',15,GRAY,align=PP_ALIGN.CENTER);txt(s,3.88,3.55,2.2,.3,'이동 중 노출',15,GRAY,align=PP_ALIGN.CENTER)
shape(s,1.49,4.15,.43,.43,BLUE,kind=MSO_SHAPE.OVAL);shape(s,4.35,4.15,.92,.43,BLUE,kind=MSO_SHAPE.OVAL);txt(s,2.86,4.08,.68,.48,'→',29,BLUE,True)
txt(s,.75,5.08,5.50,.54,'이동량 0.2 µm = 1.6 pixel',25,BLUE,True,PP_ALIGN.CENTER)
txt(s,.75,5.82,5.50,.25,'도식: 실제 영상의 크기·모양을 재현한 것은 아닙니다.',10,GRAY,align=PP_ALIGN.CENTER)
shape(s,6.80,2.80,6.07,3.48,'EAF3FF');txt(s,7.03,2.98,5.56,.42,'blur(px) = Y속도 × 노출(µs) ÷ 125',21,BLUE,True)
txt(s,7.04,3.55,5.45,.32,'X는 노출 중 정지 → X 이동속도는 blur와 분리',15,GRAY)
shape(s,7.03,4.09,5.61,.77,'DDF3EA');txt(s,7.24,4.24,5.20,.34,'허용 2px: Y100 + 노출2µs 조건 충족',19,GREEN,True)
shape(s,7.03,5.09,5.61,.77,'FFFFFF');txt(s,7.24,5.24,5.20,.34,'허용 1px: 노출2µs에서 Y ≤62.5 mm/s',19,BLUE,True)
txt(s,.50,6.55,12.3,.27,'픽셀 환산은 측정 정확도가 아닙니다. 렌즈 분해능 약1µm · 실제 밝기·SNR·검출 품질 확인이 필요합니다.',12,GRAY)
s.notes_slide.notes_text_frame.text='공식 출처: https://docs.baslerweb.com/a2a2600-20gcbas (기본2600x2128, 전체2600x2160, 2.5µm, Color, Global shutter, 기본18.5fps / Performance21.1fps). 노출2µs: https://docs.baslerweb.com/exposure-time 의 모델별 표. 렌즈: https://www.mitutoyo.com/webfoo/wp-content/uploads/E4191-378_010611.pdf p20/29. SL20x 378-810-3: NA0.28, WD30.5mm, f10mm, R1.0µm @550nm. 총배율20x는 200mm 튜브렌즈+1x relay 전제. 0.125µm/px는 샘플링 간격이며 실제 분해능·측정 정확도가 아니다. blur=100*2/125=1.6px. X 정지 촬영, Y등속. 허용1px에서는 Y≤62.5mm/s; 최소2µs보다 짧은1.25µs를 사용가능값으로 제시하지 않는다.'

# Slide 2: unique cells, actual recipe pattern, quota cycle.
s=slide(2,'같은 셀을 반복하지 않고, 헤드당 1점씩 확인합니다','고정 test0929 · 9열×5행=45 Cell · 8개 대표점의 Cell 번호 모두 다름')
shape(s,.45,1.32,6.7,5.45,'FFFFFF');txt(s,.66,1.47,6.3,.30,'실제 가공점에서 고른 서로 다른 Cell 8개',17,NAVY,True)
x0=.77;y0=2.13;sx=6.0/925;sy=3.95/850
selected={n['cell']:n for n in c['points']}
for row in range(5):
 for col in range(9):
  cid=row*9+col+1;x=x0+(15.5+col*100)*sx;y=y0+(15.5+row*200)*sy;sel=cid in selected
  shape(s,x,y,77.76*sx,46.08*sy,'D7EAFF' if sel else 'F3F6FA',BLUE if sel else 'D4DEE7',kind=MSO_SHAPE.RECTANGLE)
  txt(s,x+.005,y+.01,.5,.16,str(cid),8,BLUE if sel else GRAY,sel,PP_ALIGN.CENTER)
for n in c['points']:
 x=x0+n['x']*sx;y=y0+n['y']*sy
 shape(s,x-.03,y-.04,.07,.07,BLUE,kind=MSO_SHAPE.OVAL);txt(s,x-.10,y-.27,.70,.22,f"H{n['head']:02}",11,BLUE,True)
txt(s,.74,6.37,6.04,.22,'배치 확대도 · Cell 가로77.76 / 세로46.08 mm',11,GRAY)
for y,title,value,body in [(1.32,'2개 스캐너 / 기판','4기판 · 10분','1매 H01→H02 / 2매 H03→H04\n3매 H05→H06 / 4매 H07→H08'),(3.69,'4개 스캐너 / 기판','2기판 · 5분','1매 H01→H02→H03→H04\n2매 H05→H06→H07→H08')]:
 shape(s,7.42,y,5.45,2.14,'EAF3FF' if y<2 else 'DEF4F1');txt(s,7.64,y+.17,4.95,.32,title,18,BLUE,True);txt(s,7.64,y+.63,4.95,.48,value,29,NAVY,True);txt(s,7.64,y+1.28,4.95,.60,body,16)
txt(s,7.60,6.19,5.00,.53,'150초/기판 계획 · 기본 조건에서 두 방식 가능\n대상: 고정 대표점 8개 / 전수 검사는 별도',13,GRAY)
s.notes_slide.notes_text_frame.text='Cell 번호는 행 우선(좌→우, 위→아래). H01C01 col2 row1 B07, H02C11 col2, H03C21 col2, H04C32 col2, H05C05 col6, H06C15 col10, H07C25 col14, H08C36 col2; 전부 row1 B07. 8개 Cell은 1,11,21,32,5,15,25,36으로 중복없음. 기존 Shared src/engine.js 및 recipe.js의 유효 Shot, Mask, 소유 Head 검증통과. 물체 실제 Branch 위치, Review X와 Stage Y는 Calculation_evidence.json 참조. H01~H04 및 H05~H08 각 그룹의 ΔY는 200mm. 2개 모드 총4매, 4개 모드 총2매, 계획Tact150초/매 적용. 전체 가공점 전수 또는70 Cell/Head 조합 커버가 아니라 고정8대표점의1회 순환이다. 모든8Head 활성 전제. 원본 기하: Shared/20261002_101847_FlyingReview_Recipe_Light, 소스commit6b19d7aaf0de81271468ab7dac0fa15493172f38.'

# Slide 3: flow, worst-case timing and operational gates.
s=slide(3,'X는 먼저 도착하고, Y가 목표 위치를 지나갈 때 촬영합니다','입력 8개: 노출·허용blur / X·Y속도 / X·Y가속도·감속도 · 2개/4개 측정 선택')
steps=[('1  다른 Cell 선택','헤드당 유효 가공점 1개'),('2  X 선도착·대기','이동 → 정착 → 대기'),('3  위치 Trigger 촬영','Y 등속 · 노출 중심 보정'),('4  결과 확인·순환','Frame 수신 + Vision 결과')]
for i,(title,body) in enumerate(steps):
 x=.45+i*3.12;shape(s,x,1.38,2.88,1.08,'EAF3FF' if i!=2 else 'DEF4F1');txt(s,x+.13,1.51,2.61,.35,title,18,BLUE,True);txt(s,x+.13,1.96,2.61,.27,body,13)
shape(s,.45,2.72,8.16,3.30,'FFFFFF');txt(s,.68,2.90,7.70,.39,'가장 여유가 작은 이동: H03 → H04',22,NAVY,True)
txt(s,.68,3.43,7.60,.35,'ΔX200 mm · ΔY200 mm · X200 / Y100 mm/s',17,GRAY)
x0=.80;scale=7.14/2;barY=4.05
shape(s,x0,barY,1.2*scale,.38,CYAN,kind=MSO_SHAPE.RECTANGLE);shape(s,x0+1.2*scale,barY,.070062*scale,.38,ORANGE,kind=MSO_SHAPE.RECTANGLE);shape(s,x0+1.270062*scale,barY,.729938*scale,.38,'D7EAFF',kind=MSO_SHAPE.RECTANGLE)
txt(s,x0+.45,4.08,3.2,.27,'X 이동 1.200초',16,'FFFFFF',True,PP_ALIGN.CENTER);txt(s,5.51,4.52,2.43,.28,'대기 여유 0.730초',15,GREEN,True,PP_ALIGN.RIGHT)
ln(s,x0,4.98,x0+7.14,4.98,BLUE,3);txt(s,.79,5.12,7.20,.33,'다음 점 도착 2.000초 ≥ 준비 1.270062초',20,BLUE,True,PP_ALIGN.CENTER)
txt(s,.77,5.66,7.24,.22,'준비: 이동+정착50ms+Guard20ms+노출·지연·Jitter 여유',11,GRAY)
shape(s,8.88,2.72,3.99,3.30,'DEF4F1');txt(s,9.09,2.93,3.54,.36,'기본 조건의 판정',21,GREEN,True)
txt(s,9.10,3.50,3.49,1.03,'Blur 1.6px ≤2px\nX 선도착 여유 ≥0.730초\n촬영 모두 Y 등속 구간',17,NAVY)
txt(s,9.10,4.91,3.49,.58,'가공통과 12.232초 / 30초\n계산 Tact 100.505초 / 150초',15,GRAY)
shape(s,.45,6.27,12.42,.54,'FFFFFF');txt(s,.65,6.37,12.00,.29,'고정: 정착50ms·Guard20ms·18.5fps·150초/기판. 실제 밝기·Trigger 지연·정착성능은 장비에서 확인합니다.',12,GRAY)
s.notes_slide.notes_text_frame.text='기본값 Exposure2µs/MaxBlur2px/Vx200/Vy100/Ax=Dx1000/Ay=Dy500. 최악H03→H04 ΔX200mm, ΔY200mm; 사다리꼴 X이동1.2초, Tneed1.270062초, ΔT2초, slack0.729938초. 기본FPS18.5에서 주기54.054ms. X촬영중정지, Y계속등속; tCommand=tCross−50µs−T/2. ±5µs jitter를 시간여유에 반영. 실제 영상blur는 별도측정. Stage시작2345.9875, 끝−1584.5mm; 총거리3930.4875mm/Stroke3960. 계산가공통과12.232초는 전체유효가공첫q1845.9875~마지막q622.7875의기하시간이며, 가공전선행500mm는 전체운동에포함. Stage전체39.504875초+기타60초+영상여유1초=100.504875초. Review–Scanner거리1467mm 홀수 / 짝수1847mm(추가380)은기존명목배치가정. Scanner/Laser서비스시간, Align/왜곡, 실제Follow Error·SNR·영상PASS 미계산. 0선방어자동보정생략,실제Frame ACK/Vision 결과 이후 완료처리. 시뮬레이터에서 Blur/X timing/camera period/cruise/process/tact 조건위반시불가표시하며 요구2/4점수를자동축소하지않는다.'

out=Path('out/0선방어_검토자료.pptx');prs.save(out)
assert len(prs.slides)==3
for s in prs.slides:
 for q in s.shapes:
  assert q.left>=0 and q.top>=0 and q.left+q.width<=prs.slide_width+Inches(.02) and q.top+q.height<=prs.slide_height+Inches(.02)
print(f'Created {out}: 3 slides, editable text/shapes, all objects in slide bounds')
