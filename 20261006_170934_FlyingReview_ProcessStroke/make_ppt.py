from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pathlib import Path
import json

ev=json.loads(Path('out/Calculation_evidence.json').read_text());c=ev['fixed'];r=ev['defaultResult']
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
prs.core_properties.title='0선방어 Flying Review 검토자료';prs.core_properties.subject='Basler a2a2600-20gcBAS / Mitutoyo 378-810-3 · 서로 다른 Cell의 2/4 Head Review';prs.core_properties.author='A3 Laser Driller';prs.core_properties.keywords='Flying Vision, Zero Defence, Basler, Motion Blur'
FONT='Noto Sans CJK KR';NAVY='10243A';BLUE='1269D3';CYAN='009FA9';GRAY='5D6E80';LIGHT='F3F6FA';GREEN='0F805B';RED='BC3948';ORANGE='EFB757'
def shape(s,x,y,w,h,fill,line=None,kind=MSO_SHAPE.ROUNDED_RECTANGLE):
 q=s.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h));q.fill.solid();q.fill.fore_color.rgb=RGBColor.from_string(fill)
 if line:q.line.color.rgb=RGBColor.from_string(line);q.line.width=Pt(1)
 else:q.line.fill.background()
 if kind==MSO_SHAPE.ROUNDED_RECTANGLE:
  try:q.adjustments[0]=.10
  except:pass
 return q
def txt(s,x,y,w,h,text,size=20,color=NAVY,bold=False,align=PP_ALIGN.LEFT):
 q=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=q.text_frame;tf.clear();tf.word_wrap=True;tf.margin_left=Inches(.02);tf.margin_right=Inches(.02);tf.margin_top=Inches(.01);tf.margin_bottom=0
 for i,line in enumerate(text.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.alignment=align;p.space_before=Pt(0);p.space_after=Pt(4)
  for run in p.runs:run.font.name=FONT;run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor.from_string(color)
 return q
def ln(s,x1,y1,x2,y2,color=GRAY,width=2):
 q=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));q.line.color.rgb=RGBColor.from_string(color);q.line.width=Pt(width);return q
def slide(n,title,sub):
 s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(LIGHT)
 shape(s,0,0,13.3333,1.10,NAVY,kind=MSO_SHAPE.RECTANGLE)
 txt(s,.42,.15,12.45,.51,title,27,'FFFFFF',True);txt(s,.44,.72,12,.23,sub,11,'CFE2F4')
 txt(s,.44,7.08,11.8,.25,'A3 Laser Driller · 2026-10-06 · 명목 설계 계산 / 실제 장비·영상 검증 필요',10,GRAY)
 txt(s,12.32,7.02,.52,.30,f'{n}/3',12,GRAY,align=PP_ALIGN.RIGHT)
 return s

# Slide 1: precise scope and added stroke input.
s=slide(1,'역방향 가공 Stroke만 보고, 가감속 포함30초를 판정합니다','가공대기 → 역방향 가속·등속·감속 → Process Stroke 종료 · 정방향/전체Tact 제외')
for i,(title,body) in enumerate([('가속','설정 Y속도까지'),('등속 가공·리뷰','Y100 mm/s 기본 예시'),('감속','입력 Stroke 끝에서 정지')]):
 x=.45+i*4.2;shape(s,x,1.36,4.0,1.04,'FFFFFF');txt(s,x+.16,1.48,3.7,.32,title,21,BLUE,True);txt(s,x+.16,1.96,3.7,.27,body,15)
shape(s,.45,2.65,12.4,1.45,'EAF3FF');txt(s,.69,2.82,11.92,.34,'가공시간 = Stroke ÷ Y속도 + Y속도/(2×가속도) + Y속도/(2×감속도)',22,BLUE,True,PP_ALIGN.CENTER)
txt(s,.69,3.44,11.92,.30,'설정속도에 도달하는 사다리꼴 운동 기준 · 짧은 Stroke는 삼각형 운동으로 계산',14,GRAY,align=PP_ALIGN.CENTER)
for x,w,title,val,body,col in [(.45,5.98,'Stroke 3000 mm','30.200초 →30초 초과','3000/100 +100/1000 +100/1000',RED),(6.65,6.20,'Stroke 2980 mm (기본 예시)','30.000초 →30초 이내','가속0.200 + 등속29.600 + 감속0.200',GREEN)]:
 shape(s,x,4.37,w,1.51,'FFFFFF');txt(s,x+.18,4.51,w-.36,.30,title,18,BLUE,True);txt(s,x+.18,4.96,w-.36,.43,val,25,col,True);txt(s,x+.18,5.52,w-.36,.23,body,13,GRAY)
shape(s,.45,6.15,12.4,.60,'DEF4F1');txt(s,.66,6.30,11.96,.30,'입력9개: Stroke / 노출·허용blur / X속도·가속·감속 / Y속도·가속·감속',20,NAVY,True,PP_ALIGN.CENTER)
s.notes_slide.notes_text_frame.text='사용자선택:Stroke는가감속을포함한역방향전체구간. 가공대기 StageY1982.5mm에서출발, 입력Stroke끝에서정지. 기본2980mm는사용자actualStroke값이아니라Y100,A=D500에서30초를역산한예시. 실제Stroke입력필요. S>=v²/2a+v²/2b이면T=S/v+v/2a+v/2b. 삼각형이면T=sqrt(2S(1/a+1/b)), 설정Y속도미도달시등속가공불가. 정방향속도를100으로가정하는모델및기타60s/전체150sTact판정제거. Scanner진입전500mm선행은시작좌표에고정되며이번역방향시간에포함. 실제ProcessDone서비스지연미계산. Basler a2a2600-20gcBAS+Mitutoyo378-810-3 SL20x 사용.'

# Slide 2: bounded process window and both quotas.
s=slide(2,'리뷰2점·4점 모두, 입력 Stroke 안의 등속 구간에서 촬영','Scanner→Review1467mm · 해당 점 가공 후 촬영 · 전체 가공완료 대기 없이 병행')
shape(s,.45,1.31,12.40,2.53,'FFFFFF');x0=2.58;sc=9.58/30;xx=lambda t:x0+t*sc
for t in range(0,31,5):
 ln(s,xx(t),1.84,xx(t),3.49,'DFE8EF',1);txt(s,xx(t)-.22,1.50,.48,.22,f'{t}s',12,GRAY,align=PP_ALIGN.CENTER)
txt(s,.65,2.02,1.79,.28,'가공 Stroke',17,NAVY,True);shape(s,xx(0),1.99,30*sc,.33,BLUE,kind=MSO_SHAPE.RECTANGLE)
shape(s,xx(0),1.99,max(.06,r['profile']['t1']*sc),.33,ORANGE,kind=MSO_SHAPE.RECTANGLE);shape(s,xx(r['cruiseEnd']),1.99,max(.06,r['profile']['t3']*sc),.33,ORANGE,kind=MSO_SHAPE.RECTANGLE)
g=r['modes']['4']['groups'][0];txt(s,.65,2.86,1.79,.28,'4점 리뷰',17,NAVY,True);shape(s,xx(g['reviewStart']),2.83,(g['reviewEnd']-g['reviewStart'])*sc,.33,CYAN,kind=MSO_SHAPE.RECTANGLE)
for n in g['nodes']:
 ln(s,xx(n['cross']),2.76,xx(n['cross']),3.23,NAVY,1.5);txt(s,xx(n['cross'])-.23,3.34,.52,.22,f"H{n['head']:02}",11,NAVY,True,PP_ALIGN.CENTER)
ln(s,xx(30),1.85,xx(30),3.64,ORANGE,2);txt(s,8.40,2.40,3.47,.26,'最後노출25.932s'.replace('最後','마지막 '),13,CYAN,True,PP_ALIGN.RIGHT)
for x,w,title,value,body in [(.45,5.98,'2개 스캐너 / 기판','4기판 · 가공합계120초','1매 H01→02 / 2매 H03→04\n3매 H05→06 / 4매 H07→08'),(6.65,6.20,'4개 스캐너 / 기판','2기판 · 가공합계60초','1매 H01→02→03→04\n2매 H05→06→07→08')]:
 shape(s,x,4.11,w,1.61,'EAF3FF' if x<1 else 'DEF4F1');txt(s,x+.18,4.24,w-.36,.30,title,17,BLUE,True);txt(s,x+.18,4.67,w-.36,.37,value,23,NAVY,True);txt(s,x+.18,5.14,w-.36,.47,body,14,NAVY)
shape(s,.45,5.97,12.4,.79,'FFFFFF');txt(s,.65,6.13,11.99,.43,'최소 Stroke 2593.176 mm ≤ 입력2980 mm ≤30초 허용2980 mm\n마지막 노출±Jitter 뒤 Y 감속거리10 mm 포함 ·8개 대표점의 Cell 모두 다름',17,GREEN,True,PP_ALIGN.CENTER)
s.notes_slide.notes_text_frame.text=f'기본2980mm,Y100,A=D500,T30. 역방향시작시간0; 설정등속0.2~29.8s. 4점 H1/5 19.93175,H2/6 21.93175,H3/7 23.93175,H4/8 25.93175. 마지막노출25.931751s. 최소Stroke=2583.175mm+100*(1us+5us)+10mm=2593.1756mm. 최소Stroke의Y가속·등속도별도확인. 2점각그룹minStroke2193.1756 또는2593.1756; 모든8점1회순환은두모드모두최소2593.1756. 요약시간120/60초는역방향가공합계이고투입반송포함Tact아님. 입력2500mm는4점중3점만등속구간,요구수량축소금지,모드불가. Stroke2590mm에서H4는이동거리안이지만Y감속중이라촬영불가. 현재test0929 유효Shot범위843.2mm/명목5.265125~13.697125s이며Process전체가LaserON이라는의미아님. Frame/Vision완료ACK는별도.'

# Slide 3: blur, X readiness and input/guide synchronization.
s=slide(3,'수식 가이드도 시뮬레이터의 현재9개 입력으로 계산합니다','입력값·2/4개 모드·기판 번호 전달 · 현재 조건이 담긴 가이드HTML 저장 가능')
for i,(title,body) in enumerate([('시뮬레이터 입력','Stroke·노출·blur·X/Y 조건'),('같은 계산 엔진','30초·Stroke 내 리뷰 판정'),('수식 가이드 결과','현재 수치가 수식·도식에 반영')]):
 x=.45+i*4.2;shape(s,x,1.37,4.0,1.26,'FFFFFF');txt(s,x+.16,1.53,3.7,.32,title,20,BLUE,True);txt(s,x+.16,2.07,3.7,.33,body,15)
shape(s,.45,2.92,5.98,1.49,'EAF3FF');txt(s,.65,3.08,5.58,.31,'노출 중 이동 blur',19,BLUE,True);txt(s,.65,3.59,5.58,.40,'Y100 ×2µs ÷125 =1.6px ≤2px',23,NAVY,True)
shape(s,6.65,2.92,6.20,1.49,'DEF4F1');txt(s,6.86,3.08,5.78,.31,'X 선도착 최소 여유',19,CYAN,True);txt(s,6.86,3.59,5.78,.40,'다음 점2.000초 − 준비1.270초 =0.730초',21,NAVY,True)
shape(s,.45,4.69,12.40,.73,NAVY);txt(s,.66,4.87,11.97,.32,'30초 충족 + blur 충족 + X 준비 + 요구2/4점 모두 Stroke 내 등속 촬영',21,'FFFFFF',True,PP_ALIGN.CENTER)
shape(s,.45,5.68,12.40,1.05,'FFFFFF');txt(s,.65,5.82,11.99,.32,'가이드 보기: 현재 조건을 주소로 전달 / 같은 출처에서 열린 가이드는 입력 변경 반영',17,BLUE,True,PP_ALIGN.CENTER);txt(s,.65,6.31,11.99,.28,'현재 조건 가이드 저장: 입력을HTML에 포함 → 단독으로 열어도 동일 결과',17,NAVY,True,PP_ALIGN.CENTER)
s.notes_slide.notes_text_frame.text='동기화우선순위:URL query(9params+mode+glass), 저장한HTML내embedded snapshot, 같은출처localStorage최근입력, 기본예시. 입력누락/invalid는기본값으로조용히대체하지않고입력오류표시. BroadcastChannel및storageevent로같은출처의열린가이드갱신; file://에서는브라우저에따라채널/storage미지원가능하므로주소snapshot과단독HTML내보내기로정확한조건전달보장. 저장한가이드는고정snapshot. 시뮬레이터돌아오는링크도조건유지. 기본Blur2.5um/20x=.125um/px; v100,T2us=.2um=1.6px. X정지촬영, 최소2us기본18.5fps. Basler a2a2600-20gcBAS 공식 https://docs.baslerweb.com/a2a2600-20gcbas ; 노출 https://docs.baslerweb.com/exposure-time ; Mitutoyo378-810-3 공식 https://www.mitutoyo.com/webfoo/wp-content/uploads/E4191-378_010611.pdf p20/29. 총20x=200mm튜브렌즈+1xRelay전제. NA.28 WD30.5R약1um. X정착50msGuard20ms지연50usJitter±5us고정설계값. 실제서비스시간/S-Curve/FollowError/조명/SNR/영상검출별도검증. 0선방어자동Offset적용범위아님.'

out=Path('out/0선방어_검토자료.pptx');prs.save(out);assert len(prs.slides)==3
for s in prs.slides:
 for q in s.shapes:
  assert q.left>=0 and q.top>=0 and q.left+q.width<=prs.slide_width+Inches(.02) and q.top+q.height<=prs.slide_height+Inches(.02)
print(f'Created {out}: 3 editable slides, all shapes within slide bounds')
