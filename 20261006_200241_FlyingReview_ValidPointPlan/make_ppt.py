from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pathlib import Path
import json

ev=json.loads(Path('out/Calculation_evidence.json').read_text());c=ev['fixed'];r=ev['defaultResult']
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
prs.core_properties.title='0선방어 Flying Review 검토자료';prs.core_properties.subject='Basler a2a2600-20gcBAS / Mitutoyo 378-810-3 · 서로 다른 Cell의 2/4 Head Review';prs.core_properties.author='A3 Laser Driller';prs.core_properties.keywords='Flying Vision, Zero Defence, Basler, Motion Blur'
FONT='Noto Sans CJK KR';NAVY='10243A';BLUE='1269D3';CYAN='009FA9';GRAY='5D6E80';LIGHT='F3F6FA';GREEN='0F805B';RED='BC3948';ORANGE='EFB757'
def shape(s,x,y,w,h,fill,line=None,kind=MSO_SHAPE.ROUNDED_RECTANGLE):
 q=s.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h));q.fill.solid();q.fill.fore_color.rgb=RGBColor.from_string(fill)
 if line:q.line.color.rgb=RGBColor.from_string(line);q.line.width=Pt(1)
 else:q.line.fill.background()
 if kind==MSO_SHAPE.ROUNDED_RECTANGLE:
  try:q.adjustments[0]=.10
  except:pass
 return q
def txt(s,x,y,w,h,text,size=20,color=NAVY,bold=False,align=PP_ALIGN.LEFT):
 q=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=q.text_frame;tf.clear();tf.word_wrap=True;tf.margin_left=Inches(.02);tf.margin_right=Inches(.02);tf.margin_top=Inches(.01);tf.margin_bottom=0
 for i,line in enumerate(text.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.alignment=align;p.space_before=Pt(0);p.space_after=Pt(4)
  for run in p.runs:run.font.name=FONT;run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor.from_string(color)
 return q
def ln(s,x1,y1,x2,y2,color=GRAY,width=2):
 q=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));q.line.color.rgb=RGBColor.from_string(color);q.line.width=Pt(width);return q
def slide(n,title,sub):
 s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(LIGHT)
 shape(s,0,0,13.3333,1.10,NAVY,kind=MSO_SHAPE.RECTANGLE)
 txt(s,.42,.15,12.45,.51,title,27,'FFFFFF',True);txt(s,.44,.72,12,.23,sub,11,'CFE2F4')
 txt(s,.44,7.08,11.8,.25,'A3 Laser Driller · 2026-10-06 · 명목 설계 계산 / 실제 장비·영상 검증 필요',10,GRAY)
 txt(s,12.32,7.02,.52,.30,f'{n}/3',12,GRAY,align=PP_ALIGN.RIGHT)
 return s

# Slide 1: current conditions and four independent checks.
s=slide(1,'현재 조건과 유효 선정점으로, 30초 안에 리뷰 가능합니다','Basler a2a2600-20gcBAS + Mitutoyo378-810-3 SL20x · 가감속 포함 역방향 구간 검토')
shape(s,.45,1.36,12.4,.99,NAVY);txt(s,.67,1.53,11.96,.52,'Stroke2700mm · 노출10µs · 허용blur8px\nX200mm/s·가감속150mm/s² / Y100mm/s·가감속150mm/s²',20,'FFFFFF',True,PP_ALIGN.CENTER)
items=[('역방향30초','27.667초 ≤30초','가속0.667 + 등속26.333 + 감속0.667'),('기판·리뷰 등속','기판1500mm·후보4/4점','등속거리2633.3mm · 리뷰20.165~26.165초'),('이동 blur','8.00px = 허용8.00px','1.000µm 이동 · 허용값 경계'),('X 준비·취득','최소 여유0.211초','H03→04 여유0.275초 / H07→08 여유0.211초')]
for i,(title,value,note) in enumerate(items):
 x=.45+(i%2)*6.2;y=2.64+(i//2)*1.57
 shape(s,x,y,6.0,1.31,'FFFFFF');shape(s,x,y,.07,1.31,GREEN,kind=MSO_SHAPE.RECTANGLE)
 txt(s,x+.18,y+.11,5.61,.28,title,17,GREEN,True);txt(s,x+.18,y+.48,5.61,.36,value,23,NAVY,True);txt(s,x+.18,y+.98,5.61,.26,note,13,GRAY)
shape(s,.45,5.99,12.4,.69,'DEF4F1');txt(s,.66,6.15,11.96,.34,'8개 Head마다1점 · 모든 Cell이 다름 · 선정점의 실제 X좌표 차이로 계산',20,NAVY,True,PP_ALIGN.CENTER)
s.notes_slide.notes_text_frame.text='입력9개만 사용. Stroke2700,exp10us,blur8px,XV200,A=D150,YV100,A=D150. 시간27.666667s, Y등속0.666667~27s, 가속/감속거리각33.3333mm, 등속거리2633.3333mm. Scanner의기판1500mm통과도등속범위안. 마지막노출26.165088s±Jitter, 최소Stroke2616.509333mm. Blur8px는허용경계이며실제영상품질·측정정확도보장값아님. PointPlan H1C1,H2C11,H3C21,H4C31,H5C5,H6C15,H7C25,H8C35. 실제Frame/Vision처리완료는별도. 200mm튜브렌즈+1xRelay 총20x 전제, 물체면0.125um/px.'

# Slide 2: scanner field and actual point displacement.
s=slide(2,'가공폭110mm와, 리뷰 X 이동거리는 각각 계산합니다','H03 Cell21 X204.05mm → H04 Cell31 X306.75mm · 실제 이동102.7mm · Y차200mm')
shape(s,.45,1.34,12.4,2.48,'FFFFFF');xx=lambda z:1.46+(z-190)*.044
for h,lo,hi,fill in [(3,194.5,304.5,'EAF3FF'),(4,304.5,414.5,'DEF4F1')]:
 shape(s,xx(lo),1.65,(hi-lo)*.044,.69,fill);txt(s,xx(lo),1.81,(hi-lo)*.044,.32,f'H{h:02} 가공영역110mm',20,NAVY,True,PP_ALIGN.CENTER)
for h,xval,cell in [(3,204.05,21),(4,306.75,31)]:
 ln(s,xx(xval),2.36,xx(xval),2.78,GREEN,2);shape(s,xx(xval)-.052,2.43,.104,.104,GREEN,kind=MSO_SHAPE.OVAL);txt(s,xx(xval)-.96,2.88,1.98,.28,f'H{h:02} / C{cell}',16,GREEN,True,PP_ALIGN.CENTER)
ln(s,xx(204.05),3.39,xx(306.75),3.39,GREEN,4);txt(s,3.06,3.47,4.40,.26,'실제 ΔX=306.75−204.05=102.7mm',16,GREEN,True,PP_ALIGN.CENTER)
shape(s,.45,4.07,12.4,1.35,'EAF3FF');txt(s,.66,4.24,11.96,.34,'다음 점200/100=2.000초 / X 이동1.655 + 정착·여유0.070 =1.725초',21,NAVY,True,PP_ALIGN.CENTER)
txt(s,.66,4.82,11.96,.31,'H03→H04:2.000−1.725=여유0.275초',23,GREEN,True,PP_ALIGN.CENTER)
shape(s,.45,5.67,12.4,1.03,'FFFFFF');txt(s,.65,5.82,11.99,.34,'H08도 Cell35의 X747.25mm 유효점 선정 → H07→H08 이동110.8mm',18,NAVY,True,PP_ALIGN.CENTER)
txt(s,.65,6.29,11.99,.25,'H04:열3·행1·B07 / H08:열18·행1·B07 · Head소유권·Mask 검증 완료',14,GRAY,align=PP_ALIGN.CENTER)
s.notes_slide.notes_text_frame.text='110mm는Scanner가공폭. 실제Review이동은두선정점의ReviewX좌표차. 모델Head중심Pitch110mm이며기존좌표변환의Shot중심소유구간H3[194.5,304.5],H4[304.5,414.5]AKX. B7DOEXoffset+.3375mm, 소유권은Shot중심으로판정. H3C21col2row1beam7X204.05Q-400.675. H4C31col3row1beam7X306.75Q-600.675. H7C25col14X636.45. H8C35col18X747.25. 각가공점geometryValid=true, mask.skip=false, owner일치. H3->4 Tmove2sqrt(102.7/150)=1.654891739s, need1.724961739s, slack.275038261s. H7->8 need1.788984386s, slack.211015614s. 두구간모두설정X200까지도달하지않는삼각형운동으로계산. 시야내늦은촬영으로대체하는모델아님. 전역최단점최적화는아니며현재유효점에서충분한여유확보.'

# Slide 3: overlap and independent completion within same stroke.
s=slide(3,'가공통과와 리뷰가 겹치고, 남은 리뷰도 같은 Stroke 안에 끝납니다','기판1500mm − Scanner→Review1467mm =33mm · 등속100mm/s에서 동시통과0.330초')
shape(s,.45,1.34,12.4,2.68,'FFFFFF');x0=2.86;sc=9.35/30;xt=lambda t:x0+t*sc
for t in range(0,31,5):
 ln(s,xt(t),1.88,xt(t),3.64,'DFE8EF',1);txt(s,xt(t)-.22,1.55,.49,.24,f'{t}s',12,GRAY,align=PP_ALIGN.CENTER)
for label,y,a,b,col in [('역방향Stroke',2.00,0,r['processTime'],BLUE),('기판Scanner통과',2.67,r['scannerWindow']['start'],r['scannerWindow']['end'],BLUE),('리뷰점도착창',3.35,r['modes']['4']['groups'][0]['reviewStart'],r['modes']['4']['groups'][0]['reviewEnd'],CYAN)]:
 txt(s,.65,y,2.13,.28,label,16,NAVY,True);shape(s,xt(a),y,(b-a)*sc,.28,col,kind=MSO_SHAPE.RECTANGLE)
ln(s,xt(r['overlapStart']),2.53,xt(r['overlapStart']),3.72,ORANGE,4);ln(s,xt(r['processTime']),1.88,xt(r['processTime']),3.72,ORANGE,2)
for n in r['modes']['4']['groups'][0]['nodes']:ln(s,xt(n['cross']),3.25,xt(n['cross']),3.70,NAVY,1.4)
shape(s,.45,4.23,12.4,.96,'DEF4F1');txt(s,.65,4.39,11.99,.50,'첫 리뷰20.165초 < Scanner통과 끝20.333초\n마지막 리뷰26.165초 < Y감속 시작27.000초 < Stroke끝27.667초',19,GREEN,True,PP_ALIGN.CENTER)
for x,w,title,val in [(.45,6.0,'기판당2개 스캐너','4기판 순환 · 역방향 합계110.667초'),(6.65,6.2,'기판당4개 스캐너','2기판 순환 · 역방향 합계55.333초')]:
 shape(s,x,5.44,w,.80,'FFFFFF');txt(s,x+.17,5.52,w-.34,.23,title,15,BLUE,True);txt(s,x+.17,5.88,w-.34,.24,val,16,NAVY,True)
txt(s,.52,6.46,12.29,.32,'통과창은 실제Laser ON과 별도 · 실제Frame·Vision으로ReviewDone 판정 · 반송 포함Tact 제외',13,GRAY,align=PP_ALIGN.CENTER)
s.notes_slide.notes_text_frame.text='Stage시작1982.5mm, Scanner기판진입Q1482.5,이탈Q-17.5. 역방향V100/A=D150: Scanner기판통과5.333333~20.333333s. Review기판진입20.003333s. 두위치동시통과20.003333~20.333333s=.33s. 첫선정점노출20.165078~20.165088s이Scanner통과끝보다앞이고나머지선정점은Scanner통과이후. 마지막26.165088s, Y감속27s, Stroke종료27.666667s. 실제LaserON은가공스크립트/Scanner완료신호의증거로만판정: test0929 유효Shot창843.2mm/명목5.498458~13.930458s이므로기판통과창전체LaserON주장금지. 입력Stroke검토범위는가감속포함rest-to-rest. ProcessDone/ReviewDone서비스처리시간미계산. 2/4모드는8Head각1점순환기판수4/2, 역방향합계110.666667/55.333333s이고전체장비Tact가아님. 시뮬레이터입력9개·모드·기판번호를가이드에URL/동일출처채널로전달, 저장HTMLsnapshot지원. Basler/노출/Mitutoyo사양공식소스는HTML내링크에기재.'

out=Path('out/0선방어_검토자료.pptx');prs.save(out);assert len(prs.slides)==3
for s in prs.slides:
 for q in s.shapes:
  assert q.left>=0 and q.top>=0 and q.left+q.width<=prs.slide_width+Inches(.02) and q.top+q.height<=prs.slide_height+Inches(.02)
print(f'Created {out}: 3 editable slides, all shapes within slide bounds')
